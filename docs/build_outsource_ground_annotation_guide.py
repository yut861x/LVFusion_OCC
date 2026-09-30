from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = str(Path(__file__).resolve().parent / "120度点云监督图像地面标注规范.docx")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)

styles = doc.styles
for name, size, bold in (("Normal", 10.5, False), ("Title", 24, True),
                         ("Heading 1", 16, True), ("Heading 2", 12.5, True)):
    st = styles[name]
    st.font.name = "Microsoft YaHei"; st.font.size = Pt(size); st.font.bold = bold
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.15
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcpr.append(shd)

def borders(table):
    tblpr = table._tbl.tblPr; node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "D9D9D9"); node.append(e)
    tblpr.append(node)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False; borders(t)
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,"263A53"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9.5)
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],"F3F6F9")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def bullets(items, level=0):
    for x in items:
        p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2")
        p.add_run(x)

def numbered(items):
    for x in items: doc.add_paragraph(x, style="List Number")

title=doc.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.add_run("120度点云监督图像地面标注规范")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("委外标注执行与验收手册 版本 1.0"); r.bold=True; r.font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run("适用数据包 20260321-10-45-48    数据规模 1020帧")
doc.add_paragraph()
doc.add_heading("文档目的",1)
doc.add_paragraph("本规范用于指导外包团队利用120度 Ouster 真值点云，为图像分支制作可追溯的地面监督标签。外包人员不需要重新标注已有3D BBox。每帧必须完成点云自动预分割复核、图像投影复核、必要的人工修正和结果保存。交付结果将用于训练阶段监督图像分支；实际推理仅使用图像和33度 FOV 点云。")
doc.add_paragraph("验收结论以三维点云证据为准：确认属于路面的点标为 ground；确认高于局部地面且被物体占据的点标为 non-ground；证据不足的点标为 ignore。禁止根据图像外观猜测没有点云观测的区域。")

doc.add_heading("一 标注对象与数据目录",1)
table(["目录或文件","用途","外包操作"],[
    ("imgs/<stem>.png","3840×2160图像","查看投影结果并辅助复核"),
    ("pcds_ouster/<stem>.pcd","120度真值点云","地面标注的主要依据"),
    ("pcds/<stem>.pcd","33度部署点云","仅用于对照，不作为本任务真值来源"),
    ("labels/<stem>.txt","已有3D BBox真值","框内点自动保护，禁止重新画框"),
    ("params/cam_intrinsic.json","相机内参与畸变参数","工具自动读取"),
    ("params/lidar2cam.json","点云到相机外参","工具自动读取"),
    ("ground_config.json","自动预分割默认参数","项目负责人维护，标注员不得批量自行修改"),
], [2.0,3.0,2.0])
doc.add_paragraph("所有输入文件通过相同 stem 时间戳逐帧匹配。发现图像、Ouster点云或标定文件缺失时，应将该帧登记为异常，不得借用相邻帧标签。")

doc.add_heading("二 标签定义",1)
table(["逐点标签","编码","定义","典型对象","是否可凭图像补标"],[
    ("ground","0","有可靠局部地面参照，点位于地面表面容差带内","机动车道、可通行路面表面","不可以"),
    ("non-ground","1","有可靠地面参照，点明显高于地面，或位于已有白名单BBox内","车辆、行人、护栏、路沿、墙体、树木、路障","不可以"),
    ("ignore","2","缺少可靠地面参照、遮挡、噪点、异常低点或无法判断","远处孤立点、深坑回波、错位点、无支撑区域","不适用"),
], [1.0,.55,2.45,2.0,1.0])
doc.add_paragraph("图像稀疏监督文件采用另一套编码：0=ignore，1=ground，2=non-ground。工具自动转换，标注员不得手工修改NPZ内部数值。")

doc.add_heading("三 标注范围与优先级",1)
doc.add_paragraph("逐点标注优先级从高到低如下：")
numbered([
    "已有白名单3D BBox内的点优先标为 non-ground。BBox真值高于自动地面候选。",
    "具有可靠局部地面支撑、且明显高于地面的点标为 non-ground。",
    "位于可靠地面表面容差带内的点标为 ground。",
    "其余无法确认的点保留为 ignore。",
])
doc.add_paragraph("本任务的 non-ground 不要求区分语义类别。不属于车辆等白名单类别的护栏、路沿、墙体、树木等仍应保留为 non-ground，后续可转换为 other occupied。")

doc.add_heading("四 工具启动与默认参数",1)
doc.add_paragraph("打开数据集根目录后选择一帧，点击主窗口中的“图像地面监督”。子窗口会加载同帧图像、pcds_ouster点云、3D BBox和已有标签。首次打开时所有点为ignore；已标帧自动回放。")
table(["参数","默认值","允许的常规调整","判定作用"],[
    ("参考地面Z","-0.30 m","-0.35至-0.25 m","近场初始地面高度"),
    ("地面容差","0.15 m","0.10至0.20 m","局部地面上下容差"),
    ("最大坡度","0.04","0.02至0.06","相邻径向分箱允许的高度变化"),
    ("径向分箱","1.50 m","1.50至2.00 m","局部地面估计尺度"),
], [1.3,1.0,1.7,3.0])
doc.add_paragraph("单帧因道路条件需要超出常规范围时，标注员应截图并登记原因。不得为了增加绿色区域而随意扩大容差或坡度。")

doc.add_heading("五 标准作业流程",1)
numbered([
    "选择当前帧并打开“图像地面监督”子窗口，确认图像和投影点属于同一时刻。",
    "点击“自动预分割”，等待绿色ground、橙色non-ground、灰色ignore候选显示。",
    "先检查近场道路。确认绿色点贴合路面表面，车辆、护栏和路沿上没有大面积绿色点。",
    "再检查中远场连续性。远处零散回波不得单独延伸地面范围。",
    "检查3D BBox区域。框内车辆等障碍物点应为橙色且不可被笔刷改成ground。",
    "使用图像笔刷修正误标投影点。默认类别为non-ground；修改类别前必须确认下拉框。",
    "对无法判断、遮挡严重或投影错位区域标为ignore，不得猜测。",
    "点击“立即保存”，确认状态栏不再显示未保存。关闭并重新打开抽查自动回放。",
])

doc.add_heading("六 人工修正规则",1)
doc.add_heading("6.1 应修正为 non-ground",2)
bullets([
    "车辆、行人、骑行者和其他明显高于道路表面的目标点。",
    "护栏、隔离墩、路灯杆、标志牌、桥墩、墙体、树木和灌木点。",
    "高出路面的路沿、施工设施、遗撒物或其他占据物。",
    "自动预分割误吸入车身、车顶、护栏或立面的绿色点。",
])
doc.add_heading("6.2 应修正为 ground",2)
bullets([
    "被点云实际观测、与周围路面高度连续的机动车道路面点。",
    "因点云噪声产生少量断裂，但上下文和局部高度均能确认的路面点。",
    "BBox底部附近真实落在路面的点，只有在工具未锁定且周围地面支撑充分时才可修改。",
])
doc.add_heading("6.3 必须保留为 ignore",2)
bullets([
    "图像中可见但没有Ouster点云投影证据的像素或区域。",
    "遮挡后方空间、远处孤立点、反射异常点和明显漂浮点。",
    "低于估计路面且无法确认是下坡、沟槽还是噪声的点。",
    "图像与点云明显不同步、投影整体错位或标定异常的区域。",
])

doc.add_heading("七 典型边界案例",1)
table(["场景","正确处理","禁止操作"],[
    ("车辆遮挡路面","车辆点non-ground；遮挡后的未知区域ignore","把整块车辆下方或后方补成ground"),
    ("路沿石","高出路面的竖面和顶面优先non-ground","为追求道路连续性全部刷成ground"),
    ("护栏下方可见地面","实际地面回波ground，护栏回波non-ground","按二维区域整体统一类别"),
    ("远处稀疏点","有连续地面支撑才分类，否则ignore","用少量孤立点延伸最远地面范围"),
    ("坡道或车辆俯仰","小范围调坡度并登记，逐段检查连续性","一次提高坡度后套用所有帧"),
    ("BBox与地面候选冲突","BBox内点保持non-ground并上报明显错误框","绕过锁定将框内点改为ground"),
    ("图像边缘投影","以畸变校正后的投影点为准，错位明显则上报","凭肉眼把无点区域涂满"),
], [1.25,3.1,2.65])

doc.add_heading("八 保存格式与交付目录",1)
table(["输出目录","内容","必要字段"],[
    ("ouster_ground_labels/<stem>.npz","Ouster逐点标签，与原PCD严格同序","point_labels、class_names、pcd_sha256、point_count、proposal_parameters"),
    ("image_ground_labels/<stem>.npz","图像内最近深度投影形成的稀疏监督","uv、labels、depth、image_size、source_pcd、pcd_sha256"),
], [2.0,2.4,3.0])
bullets([
    "不得重命名、重采样或重排原始PCD点。",
    "不得直接编辑NPZ内部数组，不得使用图像软件涂改PNG作为交付标签。",
    "每个输入stem必须有两份同名输出，异常帧另附清单。",
    "交付前随机关闭并重开至少20帧，确认标签可以自动回放且无校验报错。",
])

doc.add_heading("九 质量检查标准",1)
table(["检查项","单帧要求","批次验收要求"],[
    ("时序一致性","图像和点云目标位置一致，无明显运动错位","异常帧清单完整"),
    ("BBox保护","白名单BBox内无ground误标","抽检错误率不高于0.5%"),
    ("地面准确性","ground贴合真实路面，不爬上障碍物","抽检ground精度不低于98%"),
    ("非地面完整性","路面内明显障碍点标为non-ground","抽检明显障碍漏标率不高于2%"),
    ("ignore使用","证据不足区域保持ignore，无主观填充","无大面积无依据ground"),
    ("文件完整性","两份NPZ均可加载且校验通过","1020帧全量对应，异常帧有说明"),
], [1.25,3.1,2.65])
doc.add_paragraph("建议采用三级质检：标注员自检100%，外包组长复检不少于20%，委托方终验随机抽检不少于10%。同类系统性问题一旦发现，应回退整个受影响批次，不只修复抽中帧。")

doc.add_heading("十 异常上报",1)
doc.add_paragraph("出现下列任一情况时停止当前帧标注并登记：图像或点云缺失；投影整体偏移；图像与点云时间不同步；PCD校验失败；BBox明显漂移或尺寸异常；道路高度与配置相差超过0.30 m；工具无法保存或回放。")
table(["字段","填写内容"],[
    ("frame stem","完整时间戳文件名，不含扩展名"),
    ("异常类型","缺文件、不同步、标定、BBox、点云、工具或其他"),
    ("问题描述","说明位置、现象和影响，不使用“有问题”等模糊表述"),
    ("证据","全屏截图，保留帧号、图像和投影点"),
    ("处理状态","待确认、跳过、已返修或已验收"),
], [1.6,5.4])

doc.add_heading("十一 标注员提交前检查清单",1)
for item in [
    "图像、Ouster点云和帧号一致。",
    "已执行自动预分割并检查近场、中场和远场。",
    "车辆、护栏、路沿和其他障碍物上无明显ground误标。",
    "无点云证据的图像区域没有被补标为ground。",
    "无法判断的点已设置为ignore。",
    "两份NPZ已保存，关闭重开后可以回放。",
    "异常情况已登记并附截图。",
]: doc.add_paragraph("□ " + item)

doc.add_heading("十二 委托方验收说明",1)
doc.add_paragraph("验收脚本首先检查文件数量、同名对应、NPZ字段、数组形状、标签取值、PCD点数和SHA256。随后按近场、远场、车辆密集、护栏路沿、阴影遮挡和图像边缘场景分层抽样。任何通过扩大ground范围降低ignore比例的做法均视为不合格。ignore比例不作为单独绩效指标，正确表达不确定性优先于标签数量。")

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("120度点云监督图像地面标注规范  版本 1.0")
doc.core_properties.title="120度点云监督图像地面标注规范"
doc.core_properties.subject="委外标注执行与验收手册"
doc.core_properties.author="项目组"
doc.save(OUT)
print(OUT)
