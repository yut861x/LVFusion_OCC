import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).parent
SRC = ROOT / "120度点云监督图像地面标注规范.md"
OUT = ROOT / "120度点云监督图像地面标注规范.docx"

def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def add_inline(p, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part: continue
        if part.startswith("**"):
            r=p.add_run(part[2:-2]); set_font(r,bold=True)
        elif part.startswith("`"):
            r=p.add_run(part[1:-1]); set_font(r); r.font.name="Consolas"; r.font.color.rgb=RGBColor(35,65,95)
        else: set_font(p.add_run(part))

def borders(table):
    node=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement("w:"+edge); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:color"),"D9D9D9"); node.append(e)
    table._tbl.tblPr.append(node)

lines=SRC.read_text(encoding="utf-8").splitlines()
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
for name,size,bold in (("Normal",10.5,False),("Title",24,True),("Heading 1",16,True),("Heading 2",12.5,True),("Heading 3",11,True)):
    st=doc.styles[name]; st.font.name="Microsoft YaHei"; st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
    st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
doc.styles["Normal"].paragraph_format.space_after=Pt(5); doc.styles["Normal"].paragraph_format.line_spacing=1.15
for name in ("Heading 1","Heading 2","Heading 3"):
    doc.styles[name].paragraph_format.space_before=Pt(11); doc.styles[name].paragraph_format.space_after=Pt(5)

i=0; first=True
while i<len(lines):
    s=lines[i].strip()
    if not s: i+=1; continue
    if s.startswith("```"):
        code=[]; i+=1
        while i<len(lines) and not lines[i].strip().startswith("```"): code.append(lines[i]); i+=1
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25)
        for n,line in enumerate(code):
            if n: p.add_run("\n")
            r=p.add_run(line); r.font.name="Consolas"; r.font.size=Pt(9); r.font.color.rgb=RGBColor(35,65,95)
        i+=1; continue
    if s.startswith("|") and i+1<len(lines) and re.match(r"^\|[\s:|-]+\|$",lines[i+1].strip()):
        headers=[x.strip() for x in s.strip("|").split("|")]; rows=[]; i+=2
        while i<len(lines) and lines[i].strip().startswith("|"):
            rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")]); i+=1
        t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True; borders(t)
        for j,h in enumerate(headers):
            c=t.rows[0].cells[j]; c.text=h; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"263A53"); c._tc.get_or_add_tcPr().append(shd)
            for r in c.paragraphs[0].runs: set_font(r,9,True,(255,255,255))
            c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        for ri,row in enumerate(rows):
            cells=t.add_row().cells
            for j,val in enumerate(row):
                cells[j].text=val; cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri%2:
                    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F3F6F9"); cells[j]._tc.get_or_add_tcPr().append(shd)
                for p in cells[j].paragraphs:
                    p.paragraph_format.space_after=Pt(1)
                    for r in p.runs: set_font(r,8.6)
        doc.add_paragraph(); continue
    if s.startswith("# "):
        p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_inline(p,s[2:]); first=False
    elif s.startswith("## "): doc.add_heading(s[3:],1)
    elif s.startswith("### "): doc.add_heading(s[4:],2)
    elif re.match(r"^\d+\. ",s):
        p=doc.add_paragraph(style="List Number"); add_inline(p,re.sub(r"^\d+\. ","",s))
    elif s.startswith("- [ ] "):
        p=doc.add_paragraph(); add_inline(p,"□ "+s[6:])
    elif s.startswith("- "):
        p=doc.add_paragraph(style="List Bullet"); add_inline(p,s[2:])
    else:
        p=doc.add_paragraph(); add_inline(p,s)
        if len(doc.paragraphs)<=3: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    i+=1

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("120度点云监督图像地面标注规范  版本 1.1"),9)
doc.core_properties.title="120度点云监督图像地面标注规范"
doc.core_properties.subject="委外标注与验收规范"
doc.core_properties.author="项目组"
doc.save(OUT); print(OUT)
